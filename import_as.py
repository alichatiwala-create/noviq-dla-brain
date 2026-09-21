"""
import_as.py
------------
Imports a daily AS file into staging_as, then merges it into
sources_canonical using the dedup rule from section 6 of the spec:

    Match key: NSN + CAGE + normalized Part Number
    If a row already exists (e.g. from staging_reference), just flip
    its in_as/approved flags to TRUE rather than creating a duplicate
    visible row.

Handles BOTH real formats we've seen:
  - The RAW format (as<yymmdd>.txt, bundled inside BQ's zip - this is
    what the automated nightly fetch actually receives): quoted CSV,
    no header row, 4 columns in order: NSN, CAGE, Part Number,
    Company Name. Confirmed against a real 4,267-row file.
  - The .xlsx format your colleague sent manually (in case that ever
    happens again): named headers, one dated sheet.

CONFIRMED (from the real raw file): ~0.26% of rows (11 of 4267) have
something other than a clean 13-digit NSN in the first column (e.g.
"1680LN0033448") - the same rare "LN" special-identifier pattern we
saw once in the IN file. These are skipped and counted, not silently
dropped or allowed to crash the import.

Usage:
    python import_as.py path/to/as260916.txt      # the real raw format
    python import_as.py path/to/AS.xlsx            # the manual format
"""

import sys
import os
import csv
from openpyxl import load_workbook

import db
from normalize import parse_post_date_from_filename


def normalize_part_number(raw: str) -> str:
    """Trim and uppercase so 'abc-100' and 'ABC-100' match as the same part."""
    if raw is None:
        return ""
    return str(raw).strip().upper()


def parse_file_txt(path: str):
    """
    Parse the real raw format: quoted CSV, no header, 4 columns in
    order [NSN, CAGE, Part Number, Company Name].
    """
    rows = []
    skipped_invalid_nsn = 0

    with open(path, "r", encoding="latin-1", newline="") as f:
        reader = csv.reader(f)
        for r in reader:
            if len(r) < 4:
                continue
            nsn_or_key = r[0].strip()
            if not (nsn_or_key.isdigit() and len(nsn_or_key) == 13):
                skipped_invalid_nsn += 1
                continue
            rows.append({
                "nsn": nsn_or_key,
                "cage_code": r[1].strip(),
                "part_number": r[2].strip(),
                "company_name": r[3].strip(),
            })

    if skipped_invalid_nsn:
        print(f"  Skipped {skipped_invalid_nsn} row(s) whose first column wasn't a valid "
              f"13-digit NSN (a rare special-identifier pattern, not an error).")
    return rows


def parse_file_xlsx(path: str):
    """
    Parse the manually-converted .xlsx format your colleague sent once.
    The real file has two sheets: the dated data sheet (e.g. 'as260911')
    and an empty 'Sheet1'. We read whichever sheet actually has the
    NSN/CAGE/Part Number/Company Name header, rather than hard-coding
    the sheet name (which changes daily).
    """
    wb = load_workbook(path, read_only=True, data_only=True)

    data_sheet = None
    for name in wb.sheetnames:
        ws = wb[name]
        first_row = next(ws.iter_rows(max_row=1, values_only=True), None)
        if first_row and "NSN" in [str(c).strip() if c else "" for c in first_row]:
            data_sheet = name
            break

    if data_sheet is None:
        raise ValueError(f"Couldn't find a sheet with an 'NSN' header in {path}")

    ws = wb[data_sheet]
    rows = []
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header = [str(c).strip() if c else "" for c in row]
            continue
        if not any(row):
            continue
        record = dict(zip(header, row))
        rows.append({
            "nsn": str(record.get("NSN", "")).strip(),
            "cage_code": str(record.get("Approve Source CAGE", "")).strip(),
            "part_number": str(record.get("Part Numbers", "") or "").strip(),
            "company_name": str(record.get("Approved Source Company Name", "") or "").strip(),
        })
    return rows


def parse_file(path: str):
    """Dispatches to the right parser based on the file extension."""
    if path.lower().endswith(".xlsx"):
        return parse_file_xlsx(path)
    return parse_file_txt(path)


def load_staging(conn, rows, source_file, post_date):
    cur = conn.cursor()
    total = len(rows)
    for i, row in enumerate(rows, start=1):
        cur.execute(
            """
            INSERT INTO staging_as (nsn, cage_code, part_number, company_name, source_file, post_date)
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (row["nsn"], row["cage_code"], row["part_number"], row["company_name"], source_file, post_date),
        )
        if i % 500 == 0 or i == total:
            conn.commit()
            print(f"  ...staged {i}/{total} rows")
    cur.close()


def merge_into_sources_canonical(conn, rows):
    """
    For each AS row: look for an existing canonical row on
    (nsn, cage_code, normalized part_number). If found, mark
    in_as = TRUE and approved = TRUE (and fill in company_name if we
    didn't have one). If not found, create a new canonical row with
    in_reference = FALSE (since this row came from AS, not Reference).
    """
    cur = conn.cursor()
    counts = {"new": 0, "matched_existing": 0}
    total = len(rows)

    for i, row in enumerate(rows, start=1):
        nsn = row["nsn"]
        cage_code = row["cage_code"]
        part_number = normalize_part_number(row["part_number"])
        company_name = row["company_name"] or None

        if not nsn or not cage_code:
            continue

        cur.execute(
            """
            INSERT INTO sources_canonical (nsn, cage_code, part_number, company_name, in_as, approved)
            VALUES (%s, %s, %s, %s, TRUE, TRUE)
            ON CONFLICT (nsn, cage_code, part_number) DO UPDATE
            SET in_as = TRUE,
                approved = TRUE,
                company_name = COALESCE(sources_canonical.company_name, EXCLUDED.company_name),
                updated_at = now()
            RETURNING (xmax = 0) AS was_insert
            """,
            (nsn, cage_code, part_number, company_name),
        )
        was_insert = cur.fetchone()[0]
        counts["new" if was_insert else "matched_existing"] += 1

        if i % 200 == 0 or i == total:
            conn.commit()  # save progress so far - an interruption won't lose it
            print(f"  ...processed {i}/{total} rows")

    cur.close()
    return counts


def run(path: str):
    filename = os.path.basename(path)
    # Raw .txt filenames (as260916.txt) carry the date directly. The
    # manual .xlsx ("AS.xlsx") doesn't - fall back to its dated sheet
    # name (e.g. 'as260911') in that case.
    try:
        post_date = parse_post_date_from_filename(filename, "AS")
    except ValueError:
        if filename.lower().endswith(".xlsx"):
            wb = load_workbook(path, read_only=True)
            dated_sheet = next((s for s in wb.sheetnames if s.lower().startswith("as") and any(c.isdigit() for c in s)), None)
            if dated_sheet:
                post_date = parse_post_date_from_filename(dated_sheet, "as")
            else:
                raise ValueError("Couldn't determine the post date from the filename or sheet name - pass it explicitly.")
        else:
            raise

    print(f"Parsing {filename} (post date: {post_date})...")
    rows = parse_file(path)
    print(f"  Parsed {len(rows)} rows.")

    conn = db.get_connection()
    try:
        print("Loading into staging_as...")
        load_staging(conn, rows, filename, post_date)

        print("Merging into sources_canonical (Reference/AS dedup rule)...")
        counts = merge_into_sources_canonical(conn, rows)
        print(f"  New canonical source rows: {counts['new']}")
        print(f"  Matched existing rows (flagged as approved): {counts['matched_existing']}")
    finally:
        conn.close()

    print("Done.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python import_as.py path/to/as260916.txt   (or AS.xlsx)")
        sys.exit(1)
    db.run_with_retries(run, sys.argv[1])
